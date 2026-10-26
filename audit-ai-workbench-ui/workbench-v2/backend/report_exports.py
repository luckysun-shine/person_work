import io
import os
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer


SECTION_ORDER = ("summary", "highlights", "risks", "next_plan")
DEFAULT_LABELS = {
    "summary": "总体概览",
    "highlights": "完成与进展",
    "risks": "阻塞与风险",
    "next_plan": "后续计划",
}


def _clean_text(value: object) -> str:
    return str(value or "").replace("\r\n", "\n").strip()


def _cover_config(report: dict) -> dict:
    cover = report.get("cover") or {}
    if isinstance(cover, str):
        return {}
    return cover if isinstance(cover, dict) else {}


def report_markdown(report: dict, labels: dict | None = None) -> str:
    labels = {**DEFAULT_LABELS, **(labels or {})}
    start, end = report.get("period_start") or "未设置", report.get("period_end") or report.get("period_start") or "未设置"
    cover = _cover_config(report)
    parts = [f"# {_clean_text(report.get('title'))}", "", f"报告周期：{start} 至 {end}"]
    if report.get("status"):
        parts.append(f"状态：{'定稿' if report.get('status') == 'final' else '草稿'}")
    if cover.get("enabled"):
        if cover.get("org_name"):
            parts.extend(["", f"编制单位：{_clean_text(cover.get('org_name'))}"])
        if cover.get("subtitle"):
            parts.append(f"副标题：{_clean_text(cover.get('subtitle'))}")
    for key in SECTION_ORDER:
        parts.extend(["", f"## {labels[key]}", "", _clean_text((report.get("sections") or {}).get(key)) or "暂无内容"])
    if cover.get("enabled") and cover.get("footer"):
        parts.extend(["", "---", _clean_text(cover.get("footer"))])
    return "\n".join(parts) + "\n"


def _docx_font(run, size: float, bold: bool = False, color: str | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def report_docx(report: dict, labels: dict | None = None) -> bytes:
    labels = {**DEFAULT_LABELS, **(labels or {})}
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.2); section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4); section.right_margin = Cm(2.4)
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"; normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    cover = _cover_config(report)
    start, end = report.get("period_start") or "未设置", report.get("period_end") or report.get("period_start") or "未设置"
    if cover.get("enabled"):
        for _ in range(4):
            document.add_paragraph()
        org = document.add_paragraph(); org.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _docx_font(org.add_run(_clean_text(cover.get("org_name") or "审计监察部")), 14, False, "667085")
        subtitle = document.add_paragraph(); subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _docx_font(subtitle.add_run(_clean_text(cover.get("subtitle") or "内部工作报告")), 12, False, "98A2B3")
        for _ in range(2):
            document.add_paragraph()
        cover_title = document.add_paragraph(); cover_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _docx_font(cover_title.add_run(_clean_text(report.get("title"))), 22, True, "111827")
        cover_period = document.add_paragraph(); cover_period.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _docx_font(cover_period.add_run(f"报告周期：{start} 至 {end}"), 10, color="667085")
        document.add_page_break()

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _docx_font(title.add_run(_clean_text(report.get("title"))), 20, True, "111827")
    period = document.add_paragraph()
    period.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _docx_font(period.add_run(f"报告周期：{start} 至 {end}"), 9.5, color="667085")
    document.add_paragraph()

    sections = report.get("sections") or {}
    for key in SECTION_ORDER:
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(10); heading.paragraph_format.space_after = Pt(5)
        _docx_font(heading.add_run(labels[key]), 13, True, "1D4ED8")
        content = _clean_text(sections.get(key)) or "暂无内容"
        for line in content.splitlines():
            stripped = line.strip()
            if not stripped:
                continue
            paragraph = document.add_paragraph(style="List Bullet" if re.match(r"^[-•]\s*", stripped) else None)
            paragraph.paragraph_format.space_after = Pt(4)
            _docx_font(paragraph.add_run(re.sub(r"^[-•]\s*", "", stripped)), 10.5)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_text = _clean_text(cover.get("footer")) if cover.get("enabled") and cover.get("footer") else "审计 AI 个人工作台 · 生成内容请复核后使用"
    _docx_font(footer.add_run(footer_text), 8, color="98A2B3")
    output = io.BytesIO(); document.save(output)
    return output.getvalue()


def _pdf_font_name() -> str:
    candidates = [
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "msyh.ttc",
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "simhei.ttf",
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
    ]
    for font_path in candidates:
        if font_path.exists():
            try:
                pdfmetrics.registerFont(TTFont("WorkbenchCJK", str(font_path)))
                return "WorkbenchCJK"
            except Exception:
                continue
    return "Helvetica"


def report_pdf(report: dict, labels: dict | None = None) -> bytes:
    labels = {**DEFAULT_LABELS, **(labels or {})}; font = _pdf_font_name()
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=22 * mm, leftMargin=22 * mm, topMargin=20 * mm, bottomMargin=18 * mm,
                            title=_clean_text(report.get("title")), author="审计 AI 个人工作台")
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("WorkbenchTitle", parent=styles["Title"], fontName=font, fontSize=20, leading=28, textColor=colors.HexColor("#111827"), alignment=TA_CENTER, spaceAfter=8)
    meta_style = ParagraphStyle("WorkbenchMeta", parent=styles["Normal"], fontName=font, fontSize=9, leading=14, textColor=colors.HexColor("#667085"), alignment=TA_CENTER, spaceAfter=18)
    heading_style = ParagraphStyle("WorkbenchHeading", parent=styles["Heading2"], fontName=font, fontSize=13, leading=19, textColor=colors.HexColor("#1D4ED8"), spaceBefore=10, spaceAfter=6)
    body_style = ParagraphStyle("WorkbenchBody", parent=styles["BodyText"], fontName=font, fontSize=10, leading=17, textColor=colors.HexColor("#344054"), spaceAfter=5)
    cover = _cover_config(report)
    start, end = report.get("period_start") or "未设置", report.get("period_end") or report.get("period_start") or "未设置"
    story = []
    if cover.get("enabled"):
        story.extend([
            Spacer(1, 60),
            Paragraph(_clean_text(cover.get("org_name") or "审计监察部"), meta_style),
            Paragraph(_clean_text(cover.get("subtitle") or "内部工作报告"), meta_style),
            Spacer(1, 36),
            Paragraph(_clean_text(report.get("title")), title_style),
            Paragraph(f"报告周期：{start} 至 {end}", meta_style),
            PageBreak(),
        ])
    story.extend([Paragraph(_clean_text(report.get("title")), title_style), Paragraph(f"报告周期：{start} 至 {end}", meta_style)])
    sections = report.get("sections") or {}
    for key in SECTION_ORDER:
        story.append(Paragraph(labels[key], heading_style))
        content = _clean_text(sections.get(key)) or "暂无内容"
        for line in content.splitlines():
            stripped = line.strip()
            if stripped:
                prefix = "• " if re.match(r"^[-•]\s*", stripped) else ""
                safe = re.sub(r"^[-•]\s*", "", stripped).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                story.append(Paragraph(prefix + safe, body_style))

    footer_text = _clean_text(cover.get("footer")) if cover.get("enabled") and cover.get("footer") else "审计 AI 个人工作台"

    def page_footer(canvas, document):
        canvas.saveState(); canvas.setFont(font, 8); canvas.setFillColor(colors.HexColor("#98A2B3"))
        canvas.drawCentredString(A4[0] / 2, 10 * mm, f"{footer_text}  ·  第 {document.page} 页")
        canvas.restoreState()

    doc.build(story, onFirstPage=page_footer, onLaterPages=page_footer)
    return output.getvalue()
